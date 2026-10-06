from services.csv_writer import BaseExcelWriter
from fastapi import APIRouter, UploadFile, File, HTTPException, Form, Query
from fastapi.responses import StreamingResponse
import shutil
import os
import csv
import io
from uuid import uuid4
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from celery import chain
from celery.result import AsyncResult
from core.celery_app import celery_app
from core.database import sync_engine
from tasks.catalog_tasks_v2 import ( find_catalog_diff_task, get_catalog_deleted_task, get_catalog_diff_by_label_task, 
                                    process_catalog_file_v2, sync_catalog_dictionaries, update_catalog_delete_changes,  update_catalog_save_changes,
                                      update_catalog_step_1_prepare_data, update_catalog_step_2_new_tracks,update_views_task )
from services.catalog_diff_service import (
    get_catalog_diff as get_catalog_diff_rows,
    get_catalog_deleted_tracks,
    get_processing_upload_id,
    
)
from api.deps import get_current_user, User, Depends

router = APIRouter()

STORAGE_DIR = "/app/storage"

# Те же цвета, что и в таблице на фронте (catalogDiff.css): розовый — текущие (сохранённые)
# данные, зелёный — новые (пришедшие в файле), красный — изменённые значения.
DIFF_ROW_TYPE_LABELS = {"old": "Сохранено", "new": "Пришло"}
DIFF_ROW_FILL = {
    "old": PatternFill(start_color="FFF2F2", end_color="FFF2F2", fill_type="solid"),
    "new": PatternFill(start_color="F1FBF3", end_color="F1FBF3", fill_type="solid"),
}
DIFF_CHANGED_FONT = Font(color="B3261E", bold=True)


@router.post("/upload_v2")
async def upload_catalog_v2(file: UploadFile = File(...), label_id: int = Form(...), current_user: User = Depends(get_current_user)):
    if not file.filename.endswith(('.xlsx', '.csv')):
        raise HTTPException(status_code=400, detail="Invalid file type")

    file_id = str(uuid4())
    file_ext = os.path.splitext(file.filename)[1]
    file_path = os.path.join(STORAGE_DIR, f"{file_id}{file_ext}")

    if current_user is None:
        raise HTTPException(status_code=401, detail="Unauthorized")

    with open(file_path, "wb") as buffer:
        shutil.copyfileobj(file.file, buffer)

    upload_id = str(uuid4())
    workflow = chain(
        process_catalog_file_v2.s(file_path, upload_id=upload_id, original_filename=file.filename, label_id=label_id, user_id=current_user.id),
        sync_catalog_dictionaries.s("v2", upload_id=upload_id),
    )

  
    task_result = workflow.apply_async(task_id=upload_id)

    return {
        "message": "Файл принят и поставлен в очередь на обработку (v2)",
        "task_id": upload_id,
        "filename": file.filename
    }




@router.post("/recalculate_diff")
async def recalculate_diff(file: UploadFile = File(...), 
                           label_id: int = Form(...), 
                           is_additional_data: bool = Form(True),
                           current_user: User = Depends(get_current_user)):
    if not file.filename.endswith(('.xlsx')):
        raise HTTPException(status_code=400, detail="Invalid file type")

    file_id = str(uuid4())
    file_ext = os.path.splitext(file.filename)[1]
    file_path = os.path.join(STORAGE_DIR, f"{file_id}{file_ext}")

    if current_user is None:
        raise HTTPException(status_code=401, detail="Unauthorized")
    with open(file_path, "wb") as buffer:
        shutil.copyfileobj(file.file, buffer)
    upload_id = str(uuid4())
  
    workflow = chain(
        
        process_catalog_file_v2.s(file_path, upload_id=upload_id, original_filename=file.filename, label_id=label_id, user_id=current_user.id, is_additional_data=is_additional_data), 
        update_catalog_step_1_prepare_data.s("v2", upload_id=upload_id),
        find_catalog_diff_task.s(label_id=label_id, upload_id=upload_id)
      
    )

    task_result = workflow.apply_async(task_id=upload_id)

    return {
        "message": "Файл принят и поставлен в очередь на перерасчёт diff",
        "task_id": upload_id,
        "filename": file.filename
      
    }


@router.get("/diff_result/{task_id}")
async def get_diff_result(task_id: str):
    """Опрос результата цепочки recalculate_diff по task_id последней задачи (find_catalog_diff_task)."""
    task_result = AsyncResult(task_id, app=celery_app)

    if not task_result.ready():
        return {"status": task_result.state, "ready": False}

    if task_result.failed():
        return {"status": "FAILURE", "ready": True, "error": str(task_result.result)}

    return {"status": "SUCCESS", "ready": True, "result": task_result.result}


@router.post("/save_catalog_diff")
async def save_catalog_diff(label_id: int = Form(...)):

   
    workflow = chain(
        
        update_catalog_step_2_new_tracks.s(label_id=label_id),
        update_catalog_save_changes.s(label_id=label_id)
    )

    task_result = workflow.apply_async()


    return {
        "message": "Запущена задача изменения каталога",
        "task_id": task_result.id,
        "filename": None
    }

@router.post("/delete_catalog_diff")
async def delete_catalog_diff(label_id: int = Form(...)):
   
    workflow = chain(
        update_catalog_delete_changes.s(label_id=label_id)
    )

    task_result = workflow.apply_async()


    return {
        "message": "Запущена задача удаления изменений ",
        "task_id": task_result.id,
        "filename": None
    }

@router.get("/get_catalog_diff/{label_id}")
async def get_catalog_diff(label_id: int):
   
    workflow = chain(
        get_catalog_diff_by_label_task.s(label_id)
    )

    task_result = workflow.apply_async()


    return {
        "message": "Запущена задача получения изменений ",
        "task_id": task_result.id,
        "filename": None
    }


@router.get("/get_catalog_deleted/{label_id}")
async def get_catalog_deleted(label_id: int):
   
    workflow = chain(
        get_catalog_deleted_task.s(label_id)
    )
    task_result = workflow.apply_async()


    return {
        "message": "Запущена задача получения удалённых записей ",
        "task_id": task_result.id,
        "filename": None
    }

@router.get("/export_catalog/{label_id}")
async def export_catalog(label_id: int, type: str = Query(..., pattern="^(changed|deleted)$")):
    """Синхронная выгрузка в XLSX: type=changed — изменённые треки (diff, розовый/зелёный фон + красный
    текст для изменённых полей), type=deleted — удалённые треки (розовый фон как в UI)."""
    with sync_engine.begin() as conn:
        upload_id = get_processing_upload_id(conn, label_id)

        if type == "changed":
            rows = get_catalog_diff_rows(conn, upload_id, label_id=label_id)
            filename = f"catalog_changed_label_{label_id}.xlsx"
        else:
            if not upload_id:
                raise HTTPException(status_code=400, detail=f"Для лейбла (ID: {label_id}) нет активной загрузки")
            rows = get_catalog_deleted_tracks(conn, label_id, upload_id=upload_id)
            filename = f"catalog_deleted_label_{label_id}.xlsx"

    if not rows:
        raise HTTPException(status_code=404, detail="Нет данных для выгрузки")

    output = io.BytesIO()
    original_headers = list(rows[0].keys())
    for hide_col in ["row_type", "changed_fields", "diff_type"]:
        if hide_col in original_headers:
            original_headers.remove(hide_col)

    if type == "changed":
        headers = ["Статус"] + original_headers
    else:
        headers = ["Статус"] + original_headers

    excel_writer = BaseExcelWriter(output, headers)

    if type == "changed":
        COLOR_OLD = '#FFF2F2'       # Розовый (Текущие данные/Сохранено)
        COLOR_NEW = '#F1FBF3'       # Зелёный (Новый каталог/Пришло)
        TEXT_COLOR_CHANGED = '#B3261E'  # Красный текст (изменённые поля)

        for row_dict in rows:
            row_status = str(row_dict.get("row_type", "")).upper()
            status_label = "Сохранено" if row_status == "OLD" else ("Пришло" if row_status == "NEW" else "")

            row_values = [status_label]
            for h in original_headers:
                val = row_dict.get(h)
                if isinstance(val, (list, tuple, set)):
                    row_values.append("; ".join(str(v) for v in val))
                else:
                    row_values.append("" if val is None else val)

            row_bg_color = COLOR_OLD if row_status == "OLD" else (COLOR_NEW if row_status == "NEW" else None)

            changed_fields = row_dict.get("changed_fields", []) or []
            text_colors = {
                original_headers.index(field) + 1: TEXT_COLOR_CHANGED
                for field in changed_fields if field in original_headers
            }

            excel_writer.write_formatted_row(
                row=row_values,
                row_bg_color=row_bg_color,
                text_colors=text_colors
            )
    else:
        COLOR_DELETED = '#FFF2F2'  # Розовый (Удалённые)

        for r in rows:
            status_label = "Удалён"
            row_vals = [status_label]
            for h in original_headers:
                val = r.get(h)
                if isinstance(val, (list, tuple, set)):
                    row_vals.append("; ".join(str(v) for v in val))
                else:
                    row_vals.append("" if val is None else val)

            excel_writer.write_formatted_row(row=row_vals, row_bg_color=COLOR_DELETED)

    excel_writer.close()
    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )

@router.post("/update_views")
async def update_views(label_id: int = Form(...)):
    workflow = chain(
        update_views_task.s(label_id)
    )
    task_result = workflow.apply_async()

    return {
        "message": "Запущена задача обновления представлений",
        "task_id": task_result.id,
        "filename": None
    }